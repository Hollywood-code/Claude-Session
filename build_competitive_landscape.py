from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()

for section in doc.sections:
    section.orientation = 1
    new_width, new_height = section.page_height, section.page_width
    section.page_width = new_width
    section.page_height = new_height
    section.left_margin = Inches(0.5)
    section.right_margin = Inches(0.5)
    section.top_margin = Inches(0.6)
    section.bottom_margin = Inches(0.6)

style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(10)

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
t_run = title.add_run("Competitive Landscape — AI Strategy")
t_run.bold = True
t_run.font.size = Pt(18)

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
s_run = subtitle.add_run(
    "Meridian Investor Day Prep  •  Asana, Monday.com, Smartsheet, Atlassian"
)
s_run.italic = True
s_run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

doc.add_paragraph()

headers = ["Dimension", "Asana", "Monday.com", "Smartsheet", "Atlassian"]
rows = [
    ["Public AI positioning",
     "\"Work management platform with AI built in\" — soft-agentic language (\"Asana as an agent OS\")",
     "\"Work OS, supercharged with AI\" — AI as a layer across the platform, not a separate product",
     "\"The enterprise work platform you can trust with AI\" — explicitly rejects \"agentic\" framing",
     "\"The agentic enterprise platform for software and IT teams\" — most committed to agent positioning"],
    ["Brand / product name",
     "AI Studio, Smart Workflows",
     "monday AI, monday AI Agents",
     "Smartsheet AI, AI Compliance Pack",
     "Rovo, Rovo Studio"],
    ["Pricing posture",
     "Bundled into Advanced & Enterprise tiers, no extra cost",
     "Bundled into Pro tier and above, no extra cost",
     "Base AI bundled into Business / Enterprise; Compliance Pack premium add-on ~$15/seat/mo (enterprise only)",
     "Consumption hybrid — separate paid product; ~$20/seat/mo + per-action consumption above a baseline"],
    ["Flagship announcement (last 6 mo.)",
     "Nov 2025: bundled AI Studio + Smart Workflows into Advanced / Enterprise tiers at no extra cost",
     "Jan 2026: GA of customer-buildable monday AI Agents, bundled into Pro+",
     "Dec 2025: AI Compliance Pack — agent audit logs, model selection, data residency, content filtering",
     "Jan 2026: Rovo Studio — developer-targeted agent builder running natively on Jira / Confluence / Bitbucket"],
    ["Model posture",
     "Studiously model-agnostic",
     "Multi-model, no lab named",
     "BYO-model in Compliance Pack",
     "References Anthropic by name (only one of the four)"],
    ["Governance strength",
     "Weak; lags Meridian materially",
     "Limited — no FedRAMP, minimal HIPAA, no pharma vertical",
     "Strongest of the four — governance is the explicit through-line",
     "Strong inside engineering / IT; thin outside dev orgs"],
    ["Closest to Meridian option",
     "B (Agentic) — but with bundled pricing and weak governance",
     "A in framing, executed as B-lite — broad and cheap, not deep",
     "A (PM-with-AI) — the explicit Option A peer",
     "B (Agentic platform) — the explicit Option B peer"],
    ["What it means for Meridian if we pick B",
     "Asana looks like the broader / cheaper agent; we differentiate on governance",
     "Monday is the cheaper bundled alternative; we differentiate on enterprise depth",
     "Smartsheet is our cautious cousin — we tell the \"saw the shift earlier\" story",
     "The direct comparable — analysts will ask \"why Meridian not Atlassian?\" Answer: regulated-industry breadth Atlassian doesn't serve"],
]

table = doc.add_table(rows=1 + len(rows), cols=len(headers))
table.style = "Light Grid Accent 1"
table.autofit = False

col_widths = [Inches(1.5), Inches(2.2), Inches(2.2), Inches(2.2), Inches(2.2)]
for col_idx, width in enumerate(col_widths):
    for row in table.rows:
        row.cells[col_idx].width = width

hdr_cells = table.rows[0].cells
for i, h in enumerate(headers):
    cell = hdr_cells[i]
    p = cell.paragraphs[0]
    run = p.add_run(h)
    run.bold = True
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "2c3e50")
    tc_pr.append(shd)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

for r_idx, row_data in enumerate(rows, start=1):
    cells = table.rows[r_idx].cells
    for c_idx, val in enumerate(row_data):
        cell = cells[c_idx]
        p = cell.paragraphs[0]
        run = p.add_run(val)
        run.font.size = Pt(9)
        if c_idx == 0:
            run.bold = True
        cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP

doc.add_paragraph()

q_title = doc.add_paragraph()
qt_run = q_title.add_run("The quadrant view")
qt_run.bold = True
qt_run.font.size = Pt(13)

quad_headers = ["", "PM-with-AI", "Agentic Platform"]
quad_rows = [
    ["Bundled / consumer-friendly pricing", "Monday.com", "Asana"],
    ["Premium / governance-led pricing", "Smartsheet", "Atlassian — and Meridian (proposed)"],
]

qtable = doc.add_table(rows=1 + len(quad_rows), cols=3)
qtable.style = "Light Grid Accent 1"

for i, h in enumerate(quad_headers):
    cell = qtable.rows[0].cells[i]
    p = cell.paragraphs[0]
    run = p.add_run(h)
    run.bold = True
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "2c3e50")
    tc_pr.append(shd)

for r_idx, row_data in enumerate(quad_rows, start=1):
    for c_idx, val in enumerate(row_data):
        cell = qtable.rows[r_idx].cells[c_idx]
        p = cell.paragraphs[0]
        run = p.add_run(val)
        run.font.size = Pt(10)
        if c_idx == 0:
            run.bold = True
        if "Meridian" in val:
            run.bold = True
            run.font.color.rgb = RGBColor(0xc0, 0x39, 0x2b)

doc.add_paragraph()
note = doc.add_paragraph()
n_run = note.add_run(
    "White space: no competitor today occupies the agentic-platform-with-strong-broad-governance "
    "quadrant. Atlassian is the nearest, but its governance story is dev-centric. "
    "That is the slot Meridian is proposed to claim at Investor Day."
)
n_run.italic = True
n_run.font.size = Pt(10)
n_run.font.color.rgb = RGBColor(0x44, 0x44, 0x44)

doc.save("competitive_landscape.docx")
print("Saved competitive_landscape.docx")
